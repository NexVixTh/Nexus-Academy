import json
import hashlib
import io
import sqlite3
from pathlib import Path
from typing import Callable

import httpx
import pymupdf
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app import ingestion
from app.ingestion import extract_document
from app.main import app
from app.storage import get_database_path


def create_pdf(path: Path) -> bytes:
    document = pymupdf.open()
    first_page = document.new_page()
    first_page.insert_text((72, 72), "First page")
    document.new_page()
    third_page = document.new_page()
    third_page.insert_text((72, 72), "Third page")
    document.save(path)
    document.close()
    return path.read_bytes()


def create_pptx(path: Path) -> bytes:
    presentation = Presentation()
    first_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = first_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    text_box.text_frame.text = "First slide"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)
    return path.read_bytes()


def create_rich_pptx(image_color: tuple[int, int, int] = (30, 120, 210)) -> bytes:
    image_buffer = io.BytesIO()
    Image.new("RGB", (3, 2), image_color).save(image_buffer, format="PNG")
    image_data = image_buffer.getvalue()

    presentation = Presentation()
    first_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table_shape = first_slide.shapes.add_table(
        rows=2,
        cols=2,
        left=Inches(0.5),
        top=Inches(0.5),
        width=Inches(5),
        height=Inches(1.5),
    )
    table_shape.table.cell(0, 0).text = "Component"
    table_shape.table.cell(0, 1).text = "Purpose"
    table_shape.table.cell(1, 0).text = "Ammeter"
    table_shape.table.cell(1, 1).text = "Measures circuit current"
    first_slide.shapes.add_picture(
        io.BytesIO(image_data), Inches(0.5), Inches(2.5), width=Inches(1), height=Inches(1)
    )
    first_slide.notes_slide.notes_text_frame.text = "Explain why current is measured in series."

    second_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    second_slide.shapes.add_picture(
        io.BytesIO(image_data), Inches(0.5), Inches(0.5), width=Inches(1), height=Inches(1)
    )

    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


@pytest.fixture
def database_path(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Path:
    path = tmp_path / "test-nexus.sqlite3"
    monkeypatch.setenv("NEXUS_DATABASE_PATH", str(path))
    return path


@pytest.mark.anyio
async def test_health_endpoint_remains_available() -> None:
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_pdf_extraction_preserves_page_numbers_and_empty_pages(tmp_path: Path) -> None:
    pdf_path = tmp_path / "lesson.pdf"
    create_pdf(pdf_path)

    records = extract_document(pdf_path, ".pdf")

    assert [record.source_number for record in records] == [1, 2, 3]
    assert [record.source_type for record in records] == ["page", "page", "page"]
    assert "First page" in records[0].extracted_text
    assert records[1].extracted_text == ""
    assert records[1].metadata["is_empty"] is True
    assert "Third page" in records[2].extracted_text


def test_pptx_extraction_preserves_slide_numbers_and_empty_slides(tmp_path: Path) -> None:
    pptx_path = tmp_path / "lesson.pptx"
    create_pptx(pptx_path)

    records = extract_document(pptx_path, ".pptx")

    assert [record.source_number for record in records] == [1, 2]
    assert [record.source_type for record in records] == ["slide", "slide"]
    assert records[0].extracted_text == "First slide"
    assert records[1].extracted_text == ""
    assert records[1].metadata["is_empty"] is True


def test_pptx_extracts_table_text_notes_and_raster_metadata(tmp_path: Path) -> None:
    pptx_path = tmp_path / "rich-lesson.pptx"
    pptx_path.write_bytes(create_rich_pptx())

    records = extract_document(pptx_path, ".pptx")

    assert [record.source_number for record in records] == [1, 2]
    assert "Table 1, row 1: Component | Purpose" in records[0].extracted_text
    assert "Ammeter | Measures circuit current" in records[0].extracted_text
    assert "Speaker notes: Explain why current is measured in series." in records[0].extracted_text
    assert records[0].metadata["table_count"] == 1
    assert records[0].metadata["speaker_notes_present"] is True
    assert records[0].metadata["image_count"] == 1
    assert records[1].metadata["speaker_notes_present"] is False
    assert records[1].metadata["image_count"] == 1
    assert len(records[0].assets) == len(records[1].assets) == 1

    asset = records[0].assets[0]
    assert asset.slide_number == 1
    assert asset.mime_type == "image/png"
    assert (asset.width, asset.height) == (3, 2)
    assert asset.content_hash == hashlib.sha256(asset.data).hexdigest()
    assert records[1].assets[0].slide_number == 2
    assert records[1].assets[0].content_hash == asset.content_hash


def test_oversized_raster_images_are_safely_skipped(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(ingestion, "MAX_SINGLE_IMAGE_SIZE", 1)
    pptx_path = tmp_path / "oversized-image.pptx"
    pptx_path.write_bytes(create_rich_pptx())

    records = extract_document(pptx_path, ".pptx")

    assert all(record.assets == () for record in records)
    assert [record.metadata["oversized_image_count"] for record in records] == [1, 1]


def test_unsupported_raster_images_are_safely_skipped(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(ingestion, "_sniff_raster_mime_type", lambda _: None)
    pptx_path = tmp_path / "unsupported-image.pptx"
    pptx_path.write_bytes(create_rich_pptx())

    records = extract_document(pptx_path, ".pptx")

    assert all(record.assets == () for record in records)
    assert [record.metadata["unsupported_image_count"] for record in records] == [1, 1]


@pytest.mark.anyio
async def test_upload_rejects_unsupported_extension() -> None:
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/documents", files={"file": ("notes.docx", b"not supported")}
        )

    assert response.status_code == 415


@pytest.mark.anyio
@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("broken.pdf", b"%PDF-1.7 malformed"),
        ("broken.pptx", b"not a zip package"),
    ],
)
async def test_upload_rejects_malformed_documents(filename: str, content: bytes) -> None:
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post("/documents", files={"file": (filename, content)})

    assert response.status_code == 400


@pytest.mark.anyio
async def test_upload_rejects_files_over_20_mib() -> None:
    oversized_pdf = b"%PDF-1.7" + b"x" * (20 * 1024 * 1024)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/documents", files={"file": ("large.pdf", oversized_pdf)}
        )

    assert response.status_code == 413


@pytest.mark.anyio
async def test_asset_api_deduplicates_blobs_and_scopes_slide_associations(
    database_path: Path,
) -> None:
    first_content = create_rich_pptx()
    second_content = create_rich_pptx(image_color=(210, 70, 40))
    content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        first_upload = await client.post(
            "/documents",
            files={"file": ("first.pptx", first_content, content_type)},
        )
        second_upload = await client.post(
            "/documents",
            files={"file": ("second.pptx", second_content, content_type)},
        )
        third_upload = await client.post(
            "/documents",
            files={"file": ("duplicate-content.pptx", first_content, content_type)},
        )
        assert first_upload.status_code == second_upload.status_code == third_upload.status_code == 201
        first_document_id = first_upload.json()["document_id"]
        second_document_id = second_upload.json()["document_id"]
        third_document_id = third_upload.json()["document_id"]

        first_assets_response = await client.get(
            f"/documents/{first_document_id}/assets"
        )
        second_assets_response = await client.get(
            f"/documents/{second_document_id}/assets"
        )
        third_assets_response = await client.get(
            f"/documents/{third_document_id}/assets"
        )
        missing_document_response = await client.get("/documents/missing/assets")

        first_assets = first_assets_response.json()["assets"]
        second_assets = second_assets_response.json()["assets"]
        third_assets = third_assets_response.json()["assets"]
        first_asset = first_assets[0]
        detail_response = await client.get(
            f"/documents/{first_document_id}/assets/{first_asset['asset_id']}"
        )
        cross_document_response = await client.get(
            f"/documents/{second_document_id}/assets/{first_asset['asset_id']}"
        )
        invalid_asset_id_response = await client.get(
            f"/documents/{first_document_id}/assets/not-a-uuid"
        )

    assert first_assets_response.status_code == second_assets_response.status_code == 200
    assert third_assets_response.status_code == 200
    assert [asset["slide_number"] for asset in first_assets] == [1, 2]
    assert first_assets[0]["asset_id"] == first_assets[1]["asset_id"]
    assert {asset["asset_id"] for asset in third_assets} == {first_assets[0]["asset_id"]}
    assert first_assets[0]["content_hash"] == first_assets[1]["content_hash"]
    assert first_assets[0]["mime_type"] == "image/png"
    assert (first_assets[0]["width"], first_assets[0]["height"]) == (3, 2)
    assert all("storage_key" not in asset and "path" not in asset for asset in first_assets)
    assert detail_response.status_code == 200
    assert detail_response.json()["asset_id"] == first_asset["asset_id"]
    assert cross_document_response.status_code == 404
    assert invalid_asset_id_response.status_code == 422
    assert missing_document_response.status_code == 404
    assert all(asset["asset_id"] != first_asset["asset_id"] for asset in second_assets)

    with sqlite3.connect(database_path) as connection:
        asset_count = connection.execute("SELECT COUNT(*) FROM assets").fetchone()[0]
        association_count = connection.execute(
            "SELECT COUNT(*) FROM document_assets"
        ).fetchone()[0]
        first_slide_text = connection.execute(
            "SELECT extracted_text FROM source_records "
            "WHERE document_id = ? AND source_number = 1",
            (first_document_id,),
        ).fetchone()[0]
        storage_keys = [
            row[0]
            for row in connection.execute("SELECT storage_key FROM assets").fetchall()
        ]

    assert asset_count == 2
    assert association_count == 6
    assert "Ammeter | Measures circuit current" in first_slide_text
    assert "Speaker notes: Explain why current is measured in series." in first_slide_text
    assert all((database_path.parent / "assets" / key).is_file() for key in storage_keys)


def test_asset_schema_initializes_for_an_existing_document_database(
    tmp_path: Path,
) -> None:
    database_path = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(database_path) as connection:
        connection.executescript(
            """
            CREATE TABLE documents (
                id TEXT PRIMARY KEY,
                original_filename TEXT NOT NULL,
                content_type TEXT NOT NULL,
                ingested_at TEXT NOT NULL
            );
            CREATE TABLE source_records (
                id TEXT PRIMARY KEY,
                document_id TEXT NOT NULL,
                source_type TEXT NOT NULL,
                source_number INTEGER NOT NULL,
                extracted_text TEXT NOT NULL,
                metadata_json TEXT NOT NULL
            );
            INSERT INTO documents VALUES
                ('legacy-doc', 'legacy.pptx', 'application/pptx', '2026-10-10T00:00:00+00:00');
            """
        )

    from app.storage import get_document_assets

    assert get_document_assets(database_path, "legacy-doc") == []
    with sqlite3.connect(database_path) as connection:
        existing_document = connection.execute(
            "SELECT original_filename FROM documents WHERE id = ?",
            ("legacy-doc",),
        ).fetchone()
        asset_table = connection.execute(
            "SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'assets'"
        ).fetchone()
        association_table = connection.execute(
            "SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'document_assets'"
        ).fetchone()
    assert existing_document == ("legacy.pptx",)
    assert asset_table == ("assets",)
    assert association_table == ("document_assets",)


@pytest.mark.parametrize(
    ("filename", "create_document", "source_type", "expected_records"),
    [
        (
            "course-notes.pdf",
            create_pdf,
            "page",
            [("First page", False), ("", True), ("Third page", False)],
        ),
        (
            "course-slides.pptx",
            create_pptx,
            "slide",
            [("First slide", False), ("", True)],
        ),
    ],
)
@pytest.mark.anyio
async def test_upload_persists_document_and_source_records(
    database_path: Path,
    tmp_path: Path,
    filename: str,
    create_document: Callable[[Path], bytes],
    source_type: str,
    expected_records: list[tuple[str, bool]],
) -> None:
    assert get_database_path() == database_path
    document_content = create_document(tmp_path / filename)
    content_type = (
        "application/pdf"
        if filename.endswith(".pdf")
        else "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/documents",
            files={"file": (filename, document_content, content_type)},
        )

    assert response.status_code == 201
    result = response.json()
    assert result["pages_or_slides_processed"] == len(expected_records)
    assert result["records_created"] == len(expected_records)
    assert "extracted_text" not in result

    with sqlite3.connect(database_path) as connection:
        document = connection.execute(
            "SELECT id, original_filename, content_type, ingested_at "
            "FROM documents WHERE id = ?",
            (result["document_id"],),
        ).fetchone()
        records = connection.execute(
            "SELECT id, document_id, source_type, source_number, extracted_text, metadata_json "
            "FROM source_records WHERE document_id = ? ORDER BY source_number",
            (result["document_id"],),
        ).fetchall()

    assert document is not None
    assert document[0] == result["document_id"]
    assert document[1:3] == (filename, content_type)
    assert document[3]
    assert len(records) == len(expected_records)
    assert len({record[0] for record in records}) == len(expected_records)
    assert all(record[0] for record in records)
    assert all(record[1] == result["document_id"] for record in records)
    assert [(record[2], record[3]) for record in records] == [
        (source_type, source_number)
        for source_number in range(1, len(expected_records) + 1)
    ]
    assert [record[4].strip() for record in records] == [
        expected_text for expected_text, _ in expected_records
    ]
    assert [json.loads(record[5])["is_empty"] for record in records] == [
        is_empty for _, is_empty in expected_records
    ]