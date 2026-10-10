import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Literal
from zipfile import BadZipFile, ZipFile

import pymupdf
from lxml.etree import XMLSyntaxError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError


SourceType = Literal["page", "slide"]
MAX_SINGLE_IMAGE_SIZE = 10 * 1024 * 1024
MAX_DOCUMENT_IMAGE_SIZE = 20 * 1024 * 1024
MAX_IMAGES_PER_DOCUMENT = 100
MAX_IMAGE_PIXELS = 100_000_000

_RASTER_SIGNATURES: tuple[tuple[bytes, str], ...] = (
    (b"\x89PNG\r\n\x1a\n", "image/png"),
    (b"\xff\xd8\xff", "image/jpeg"),
    (b"GIF87a", "image/gif"),
    (b"GIF89a", "image/gif"),
    (b"BM", "image/bmp"),
    (b"II*\x00", "image/tiff"),
    (b"MM\x00*", "image/tiff"),
)
_RASTER_EXTENSIONS = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
    "image/tiff": ".tif",
}


class UnsupportedDocumentError(ValueError):
    pass


class InvalidDocumentError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class ExtractedAsset:
    slide_number: int
    data: bytes
    mime_type: str
    width: int
    height: int
    content_hash: str


@dataclass(frozen=True, slots=True)
class ExtractedRecord:
    source_type: SourceType
    source_number: int
    extracted_text: str
    metadata: dict[str, bool | int]
    assets: tuple[ExtractedAsset, ...] = ()


def extract_document(path: Path, extension: str) -> list[ExtractedRecord]:
    normalized_extension = extension.lower()
    if normalized_extension == ".pdf":
        return _extract_pdf(path)
    if normalized_extension == ".pptx":
        return _extract_pptx(path)
    raise UnsupportedDocumentError("Only PDF and PPTX documents are supported.")


def _extract_pdf(path: Path) -> list[ExtractedRecord]:
    with path.open("rb") as document_file:
        if document_file.read(5) != b"%PDF-":
            raise InvalidDocumentError("The uploaded file is not a valid PDF.")

    try:
        with pymupdf.open(stream=path.read_bytes(), filetype="pdf") as document:
            if document.is_encrypted:
                raise InvalidDocumentError("Password-protected PDFs are not supported.")
            if document.page_count == 0:
                raise InvalidDocumentError("The PDF contains no pages.")

            records = []
            for page_index in range(document.page_count):
                text = document.load_page(page_index).get_text("text") or ""
                records.append(
                    ExtractedRecord(
                        source_type="page",
                        source_number=page_index + 1,
                        extracted_text=text,
                        metadata={"is_empty": not text.strip(), "text_length": len(text)},
                    )
                )
            return records
    except InvalidDocumentError:
        raise
    except (pymupdf.FileDataError, OSError, RuntimeError) as exc:
        raise InvalidDocumentError("The PDF file is malformed or unreadable.") from exc


def _extract_pptx(path: Path) -> list[ExtractedRecord]:
    try:
        with ZipFile(path) as package:
            package_parts = set(package.namelist())
            if not {"[Content_Types].xml", "ppt/presentation.xml"}.issubset(package_parts):
                raise InvalidDocumentError("The uploaded file is not a valid PPTX presentation.")

        presentation = Presentation(path)
        if not presentation.slides:
            raise InvalidDocumentError("The PPTX contains no slides.")

        records = []
        total_unique_image_bytes = 0
        total_image_references = 0
        seen_image_hashes: set[str] = set()
        for slide_number, slide in enumerate(presentation.slides, start=1):
            text_parts: list[str] = []
            table_count = 0
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    text_parts.append(shape.text)
                if shape.has_table:
                    table_count += 1
                    for row_number, row in enumerate(shape.table.rows, start=1):
                        cells = [" ".join(cell.text.split()) for cell in row.cells]
                        if any(cells):
                            text_parts.append(
                                f"Table {table_count}, row {row_number}: "
                                + " | ".join(cells)
                            )

            notes_present = False
            try:
                if slide.has_notes_slide:
                    notes_frame = slide.notes_slide.notes_text_frame
                    notes_text = notes_frame.text.strip() if notes_frame is not None else ""
                    if notes_text:
                        notes_present = True
                        text_parts.append(f"Speaker notes: {notes_text}")
            except Exception:
                notes_present = False

            slide_assets: list[ExtractedAsset] = []
            unsupported_image_count = 0
            oversized_image_count = 0
            malformed_image_count = 0
            for picture in _iter_picture_shapes(slide.shapes):
                if total_image_references >= MAX_IMAGES_PER_DOCUMENT:
                    oversized_image_count += 1
                    continue
                try:
                    image = picture.image
                    blob = image.blob
                    if not blob:
                        malformed_image_count += 1
                        continue
                    if len(blob) > MAX_SINGLE_IMAGE_SIZE:
                        oversized_image_count += 1
                        continue
                    mime_type = _sniff_raster_mime_type(blob)
                    if mime_type is None:
                        unsupported_image_count += 1
                        continue
                    width, height = image.size
                    if (
                        width <= 0
                        or height <= 0
                        or width * height > MAX_IMAGE_PIXELS
                    ):
                        oversized_image_count += 1
                        continue

                    content_hash = hashlib.sha256(blob).hexdigest()
                    if (
                        content_hash not in seen_image_hashes
                        and total_unique_image_bytes + len(blob) > MAX_DOCUMENT_IMAGE_SIZE
                    ):
                        oversized_image_count += 1
                        continue
                    if content_hash not in seen_image_hashes:
                        total_unique_image_bytes += len(blob)
                        seen_image_hashes.add(content_hash)
                    slide_assets.append(
                        ExtractedAsset(
                            slide_number=slide_number,
                            data=blob,
                            mime_type=mime_type,
                            width=width,
                            height=height,
                            content_hash=content_hash,
                        )
                    )
                    total_image_references += 1
                except Exception:
                    malformed_image_count += 1

            text = "\n".join(text_parts)
            records.append(
                ExtractedRecord(
                    source_type="slide",
                    source_number=slide_number,
                    extracted_text=text,
                    metadata={
                        "is_empty": not text.strip(),
                        "text_length": len(text),
                        "table_count": table_count,
                        "speaker_notes_present": notes_present,
                        "image_count": len(slide_assets),
                        "unsupported_image_count": unsupported_image_count,
                        "oversized_image_count": oversized_image_count,
                        "malformed_image_count": malformed_image_count,
                    },
                    assets=tuple(slide_assets),
                )
            )
        return records
    except InvalidDocumentError:
        raise
    except (BadZipFile, KeyError, OSError, PackageNotFoundError, ValueError, XMLSyntaxError) as exc:
        raise InvalidDocumentError("The PPTX file is malformed or unreadable.") from exc


def _iter_picture_shapes(shapes: object) -> list[object]:
    pictures: list[object] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pictures.append(shape)
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            pictures.extend(_iter_picture_shapes(shape.shapes))
    return pictures


def _sniff_raster_mime_type(data: bytes) -> str | None:
    return next(
        (mime_type for signature, mime_type in _RASTER_SIGNATURES if data.startswith(signature)),
        None,
    )