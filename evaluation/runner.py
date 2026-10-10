from __future__ import annotations

import argparse
import asyncio
import io
import json
import os
import re
import sys
import tempfile
import unicodedata
from pathlib import Path
from typing import Any

import httpx
import pymupdf
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.main import app  # noqa: E402
from app.quiz import question_concept  # noqa: E402
from app.storage import StoredSourceRecord, get_database_path, get_document_source_records  # noqa: E402


DEFAULT_DATASET = ROOT / "evaluation" / "fixtures" / "task7_v1.json"
DEFAULT_REPORT_DIRECTORY = ROOT / "evaluation" / "reports"
_URL_ONLY_PATTERN = re.compile(r"^(?:https?://|www\.)\S+$", re.IGNORECASE)
_GENERIC_HEADINGS = {
    "url:",
    "qr code:",
    "algorithm explanation",
    "naive method",
}
_SOURCE_LABELS = {"page": "pdf_page", "slide": "pptx_slide"}


class BenchmarkDatasetError(ValueError):
    pass


def load_dataset(path: Path = DEFAULT_DATASET) -> dict[str, Any]:
    try:
        raw_data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise BenchmarkDatasetError("Benchmark dataset could not be read as JSON.") from exc
    return validate_dataset(raw_data)


def validate_dataset(data: object) -> dict[str, Any]:
    if not isinstance(data, dict):
        raise BenchmarkDatasetError("Dataset root must be an object.")
    if not isinstance(data.get("dataset_version"), str) or not data["dataset_version"].strip():
        raise BenchmarkDatasetError("Dataset must declare a non-empty dataset_version.")

    documents = data.get("documents")
    if not isinstance(documents, list) or not documents:
        raise BenchmarkDatasetError("Dataset must contain at least one document.")
    documents_by_id: dict[str, dict[str, Any]] = {}
    for document in documents:
        if not isinstance(document, dict):
            raise BenchmarkDatasetError("Each document must be an object.")
        document_id = document.get("id")
        document_format = document.get("format")
        records = document.get("records")
        if not isinstance(document_id, str) or not document_id:
            raise BenchmarkDatasetError("Each document needs a non-empty id.")
        if document_id in documents_by_id:
            raise BenchmarkDatasetError("Document IDs must be unique.")
        if document_format not in {"pdf", "pptx"}:
            raise BenchmarkDatasetError("Document format must be pdf or pptx.")
        if not isinstance(records, list) or not records:
            raise BenchmarkDatasetError("Each document needs at least one source record.")
        locations: set[int] = set()
        for record in records:
            if not isinstance(record, dict):
                raise BenchmarkDatasetError("Each source record must be an object.")
            location = record.get("location")
            text = record.get("text")
            if not isinstance(location, int) or isinstance(location, bool) or location < 1:
                raise BenchmarkDatasetError("Source locations must be positive integers.")
            if location in locations:
                raise BenchmarkDatasetError("Source locations must be unique per document.")
            if not isinstance(text, str) or not text.strip():
                raise BenchmarkDatasetError("Source record text must be non-empty.")
            locations.add(location)
        documents_by_id[document_id] = document

    qa_cases = data.get("qa_cases")
    if not isinstance(qa_cases, list) or not qa_cases:
        raise BenchmarkDatasetError("Dataset must contain at least one QA case.")
    seen_case_ids: set[str] = set()
    for case in qa_cases:
        if not isinstance(case, dict):
            raise BenchmarkDatasetError("Each QA case must be an object.")
        case_id = case.get("id")
        document = documents_by_id.get(case.get("document_id"))
        if not isinstance(case_id, str) or not case_id or case_id in seen_case_ids:
            raise BenchmarkDatasetError("QA case IDs must be non-empty and unique.")
        seen_case_ids.add(case_id)
        if document is None:
            raise BenchmarkDatasetError("QA case references an unknown document.")
        if not isinstance(case.get("question"), str) or not case["question"].strip():
            raise BenchmarkDatasetError("QA case question must be non-empty.")
        if not isinstance(case.get("answerable"), bool):
            raise BenchmarkDatasetError("QA case must declare answerable as a boolean.")
        terms = case.get("expected_answer_terms")
        if not isinstance(terms, list) or any(not isinstance(term, str) for term in terms):
            raise BenchmarkDatasetError("expected_answer_terms must be a list of strings.")
        expected_source = case.get("expected_source")
        if case["answerable"]:
            if not isinstance(expected_source, dict):
                raise BenchmarkDatasetError("Answerable cases need an expected source location.")
            expected_type = "pdf_page" if document["format"] == "pdf" else "pptx_slide"
            if expected_source.get("type") != expected_type:
                raise BenchmarkDatasetError("Expected source type does not match document format.")
            if not any(
                record["location"] == expected_source.get("location")
                for record in document["records"]
            ):
                raise BenchmarkDatasetError("QA case references an absent source location.")
            if not terms:
                raise BenchmarkDatasetError("Answerable cases need manually verified answer terms.")
        elif expected_source is not None or terms:
            raise BenchmarkDatasetError(
                "Unanswerable cases cannot declare expected citations or answer terms."
            )

    quiz_checks = data.get("quiz_checks")
    if not isinstance(quiz_checks, list) or not quiz_checks:
        raise BenchmarkDatasetError("Dataset must contain quiz checks.")
    for check in quiz_checks:
        if not isinstance(check, dict) or check.get("document_id") not in documents_by_id:
            raise BenchmarkDatasetError("Quiz check references an unknown document.")
        if not isinstance(check.get("requested_count"), int) or not 1 <= check["requested_count"] <= 10:
            raise BenchmarkDatasetError("Quiz requested_count must be between 1 and 10.")
        blocked = check.get("must_reject_candidates")
        if not isinstance(blocked, list) or any(not isinstance(text, str) for text in blocked):
            raise BenchmarkDatasetError("must_reject_candidates must be a list of strings.")

    learner_scenario = data.get("learner_scenario")
    if not isinstance(learner_scenario, dict) or any(
        learner_scenario.get(key) not in documents_by_id
        for key in ("document_id", "isolation_document_id")
    ):
        raise BenchmarkDatasetError("Learner scenario must reference known documents.")
    if not isinstance(learner_scenario.get("correct_repetitions_for_strong_concept"), int):
        raise BenchmarkDatasetError("Learner scenario needs an integer repetition count.")

    return data


def rate(numerator: int, denominator: int) -> dict[str, int | float | None]:
    return {
        "numerator": numerator,
        "denominator": denominator,
        "value": round(numerator / denominator, 4) if denominator else None,
    }


def make_pdf_bytes(records: list[dict[str, Any]]) -> bytes:
    document = pymupdf.open()
    for record in records:
        page = document.new_page()
        page.insert_textbox(
            pymupdf.Rect(54, 54, 558, 738),
            record["text"],
            fontsize=11,
        )
    content = document.tobytes()
    document.close()
    return content


def make_pptx_bytes(records: list[dict[str, Any]]) -> bytes:
    presentation = Presentation()
    for record in records:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        text_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(6))
        text_box.text_frame.text = record["text"]
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


async def upload_fixture_documents(
    client: httpx.AsyncClient,
    dataset: dict[str, Any],
) -> tuple[dict[str, str], dict[str, dict[str, StoredSourceRecord]]]:
    document_ids: dict[str, str] = {}
    records_by_document: dict[str, dict[str, StoredSourceRecord]] = {}
    for document in dataset["documents"]:
        extension = document["format"]
        filename = f"{document['id']}.{extension}"
        content = (
            make_pdf_bytes(document["records"])
            if extension == "pdf"
            else make_pptx_bytes(document["records"])
        )
        content_type = (
            "application/pdf"
            if extension == "pdf"
            else "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
        response = await client.post(
            "/documents",
            files={"file": (filename, content, content_type)},
        )
        if response.status_code != 201:
            raise RuntimeError("Synthetic benchmark document upload failed.")
        document_id = response.json()["document_id"]
        document_ids[document["id"]] = document_id
        stored_records = get_document_source_records(get_database_path(), document_id)
        if stored_records is None:
            raise RuntimeError("Uploaded benchmark source records could not be read back.")
        records_by_document[document["id"]] = {
            record.id: record for record in stored_records
        }
    return document_ids, records_by_document


def evaluate_qa_case(
    case: dict[str, Any],
    response_data: dict[str, Any] | None,
    records_by_id: dict[str, StoredSourceRecord],
) -> dict[str, bool]:
    returned = response_data is not None
    abstained = bool(response_data and response_data.get("abstained"))
    citations = response_data.get("citations", []) if response_data else []
    answer = str(response_data.get("answer", "")).casefold() if response_data else ""
    if not case["answerable"]:
        return {
            "abstained": abstained,
            "unsupported_answer": returned and not abstained,
        }

    expected = case["expected_source"]
    expected_internal_type = "page" if expected["type"] == "pdf_page" else "slide"
    expected_record = next(
        (
            record
            for record in records_by_id.values()
            if record.source_type == expected_internal_type
            and record.source_number == expected["location"]
        ),
        None,
    )
    citation_source_correct = any(
        expected_record is not None
        and citation.get("source_record_id") == expected_record.id
        and citation.get("source_record_id") in records_by_id
        for citation in citations
    )
    citation_location_correct = any(
        citation.get("source_type") == expected["type"]
        and citation.get("page_or_slide_number") == expected["location"]
        for citation in citations
    )
    evidence_supported = all(term.casefold() in answer for term in case["expected_answer_terms"])
    return {
        "answerable_covered": returned and not abstained,
        "citation_source_record_correct": citation_source_correct,
        "citation_location_correct": citation_location_correct,
        "citation_correct": citation_source_correct and citation_location_correct,
        "evidence_supported": evidence_supported,
        "unsupported_answer": False,
    }


async def evaluate_qa(
    client: httpx.AsyncClient,
    dataset: dict[str, Any],
    document_ids: dict[str, str],
    records_by_document: dict[str, dict[str, StoredSourceRecord]],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    outcomes: list[dict[str, Any]] = []
    for case in dataset["qa_cases"]:
        response = await client.post(
            f"/documents/{document_ids[case['document_id']]}/ask",
            json={"question": case["question"]},
        )
        response_data = response.json() if response.status_code == 200 else None
        checks = evaluate_qa_case(
            case,
            response_data,
            records_by_document[case["document_id"]],
        )
        outcomes.append({"case_id": case["id"], **checks})

    answerable = [case for case in dataset["qa_cases"] if case["answerable"]]
    unanswerable = [case for case in dataset["qa_cases"] if not case["answerable"]]
    outcome_by_id = {outcome["case_id"]: outcome for outcome in outcomes}
    fields = {
        "answerable_coverage": "answerable_covered",
        "citation_source_record_accuracy": "citation_source_record_correct",
        "citation_location_accuracy": "citation_location_correct",
        "citation_correctness": "citation_correct",
        "evidence_support_rate": "evidence_supported",
    }
    metrics = {
        name: rate(
            sum(outcome_by_id[case["id"]].get(field, False) for case in answerable),
            len(answerable),
        )
        for name, field in fields.items()
    }
    metrics["appropriate_abstention_rate"] = rate(
        sum(outcome_by_id[case["id"]]["abstained"] for case in unanswerable),
        len(unanswerable),
    )
    metrics["unsupported_answer_rate"] = rate(
        sum(outcome_by_id[case["id"]]["unsupported_answer"] for case in unanswerable),
        len(unanswerable),
    )
    return metrics, outcomes


def _source_validity(
    question: dict[str, Any],
    document_id: str,
    records_by_id: dict[str, StoredSourceRecord],
) -> bool:
    record = records_by_id.get(question.get("source_record_id"))
    if record is None:
        return False
    expected_source_type = {"page": "pdf_page", "slide": "pptx_slide"}.get(
        record.source_type
    )
    return (
        record.document_id == document_id
        and question.get("source_type") == expected_source_type
        and question.get("page_or_slide_number") == record.source_number
        and question.get("supporting_excerpt") in record.extracted_text
        and all(option.get("text") in record.extracted_text for option in question.get("options", []))
    )


def _question_quality(
    question: dict[str, Any],
    source_valid: bool,
    blocked_candidates: set[str],
) -> tuple[bool, bool]:
    options = question.get("options", [])
    option_ids = [option.get("option_id") for option in options]
    texts = [option.get("text", "") for option in options]
    answer_key = question.get("answer_key")
    correct = [option for option in options if option.get("option_id") == answer_key]
    term = question_concept(question.get("question_text", ""))
    key_consistent = bool(
        term
        and len(correct) == 1
        and term.casefold() in correct[0].get("text", "").casefold()
        and all(
            term.casefold() not in option.get("text", "").casefold()
            for option in options
            if option.get("option_id") != answer_key
        )
    )
    valid = (
        source_valid
        and len(options) == 4
        and set(option_ids) == {"A", "B", "C", "D"}
        and len(set(text.casefold() for text in texts)) == 4
        and key_consistent
        and not any(_URL_ONLY_PATTERN.fullmatch(text.strip()) for text in texts)
        and not any(text.strip().casefold() in blocked_candidates for text in texts)
    )
    return valid, key_consistent


async def evaluate_quizzes(
    client: httpx.AsyncClient,
    dataset: dict[str, Any],
    document_ids: dict[str, str],
    records_by_document: dict[str, dict[str, StoredSourceRecord]],
) -> dict[str, Any]:
    question_count_requested = 0
    question_count_generated = 0
    valid_source_count = 0
    consistent_answer_keys = 0
    quality_questions = 0
    duplicate_occurrences = 0
    blocked_options = 0
    total_options = 0
    blocked_candidates_expected = 0
    blocked_candidates_rejected = 0
    insufficient_expected = 0
    insufficient_correct = 0
    seen_questions: set[tuple[str, str]] = set()
    response_summaries: list[dict[str, Any]] = []

    for check in dataset["quiz_checks"]:
        requested = check["requested_count"]
        question_count_requested += requested
        response = await client.post(
            f"/documents/{document_ids[check['document_id']]}/quiz",
            json={"question_count": requested},
        )
        if response.status_code != 200:
            response_summaries.append({"check_id": check["id"], "status_code": response.status_code})
            continue

        data = response.json()
        questions = data["questions"]
        question_count_generated += len(questions)
        blocked = {candidate.casefold() for candidate in check["must_reject_candidates"]}
        blocked_candidates_expected += len(blocked)
        emitted_texts = {
            option.get("text", "").casefold()
            for question in questions
            for option in question.get("options", [])
        }
        blocked_candidates_rejected += sum(candidate not in emitted_texts for candidate in blocked)
        if check.get("expect_insufficient_content", False):
            insufficient_expected += 1
            insufficient_correct += int(
                len(questions) < requested and bool(data.get("limitation"))
            )

        valid_count = 0
        for question in questions:
            source_valid = _source_validity(
                question,
                document_ids[check["document_id"]],
                records_by_document[check["document_id"]],
            )
            valid_count += int(source_valid)
            quality, key_consistent = _question_quality(question, source_valid, blocked)
            quality_questions += int(quality)
            consistent_answer_keys += int(key_consistent)
            options = question.get("options", [])
            total_options += len(options)
            blocked_options += sum(
                int(
                    bool(_URL_ONLY_PATTERN.fullmatch(option.get("text", "").strip()))
                    or option.get("text", "").strip().casefold() in blocked
                )
                for option in options
            )
            identity = (check["document_id"], question.get("question_id", ""))
            duplicate_occurrences += int(identity in seen_questions)
            seen_questions.add(identity)
        valid_source_count += valid_count
        response_summaries.append(
            {
                "check_id": check["id"],
                "status_code": response.status_code,
                "requested": requested,
                "generated": len(questions),
                "limitation_reported": bool(data.get("limitation")),
            }
        )

    metrics = {
        "generated_questions_vs_requested": rate(question_count_generated, question_count_requested),
        "valid_source_reference_rate": rate(valid_source_count, question_count_generated),
        "answer_key_consistency_rate": rate(consistent_answer_keys, question_count_generated),
        "duplicate_question_rate": rate(duplicate_occurrences, question_count_generated),
        "url_or_heading_option_rate": rate(blocked_options, total_options),
        "quality_rule_satisfaction_rate": rate(quality_questions, question_count_generated),
        "blocked_candidate_rejection_rate": rate(
            blocked_candidates_rejected, blocked_candidates_expected
        ),
        "insufficient_content_report_rate": rate(insufficient_correct, insufficient_expected),
    }
    return {"metrics": metrics, "quiz_checks": response_summaries}


async def evaluate_personalization(
    client: httpx.AsyncClient,
    dataset: dict[str, Any],
    document_ids: dict[str, str],
    records_by_document: dict[str, dict[str, StoredSourceRecord]],
) -> dict[str, Any]:
    scenario = dataset["learner_scenario"]
    document_id = document_ids[scenario["document_id"]]
    learner_response = await client.post("/learners")
    if learner_response.status_code != 201:
        raise RuntimeError("Could not create simulated benchmark learner.")
    learner_id = learner_response.json()["learner_id"]

    base_response = await client.post(
        f"/documents/{document_id}/quiz",
        json={"question_count": 10, "learner_id": learner_id},
    )
    if base_response.status_code != 200 or len(base_response.json()["questions"]) < 3:
        raise RuntimeError("Learner fixture did not produce enough benchmark questions.")
    base_quiz = base_response.json()
    base_questions = base_quiz["questions"]
    weak_question = base_questions[0]
    strong_question = next(
        question
        for question in base_questions
        if question["concept_label"] != weak_question["concept_label"]
    )
    unseen_label = next(
        question["concept_label"]
        for question in base_questions
        if question["concept_label"] not in {
            weak_question["concept_label"],
            strong_question["concept_label"],
        }
    )

    wrong_option_id = next(
        option["option_id"]
        for option in weak_question["options"]
        if option["option_id"] != weak_question["answer_key"]
    )
    weak_answer = [{
        "question_id": weak_question["question_id"],
        "selected_option_id": wrong_option_id,
    }]
    weak_submit_path = f"/documents/{document_id}/quiz/submit"
    weak_request = {
        "quiz_id": base_quiz["quiz_id"],
        "learner_id": learner_id,
        "answers": weak_answer,
    }
    first_weak_submit = await client.post(weak_submit_path, json=weak_request)
    progress_after_weak = await client.get(f"/learners/{learner_id}/progress")
    retry_weak_submit = await client.post(weak_submit_path, json=weak_request)
    progress_after_retry = await client.get(f"/learners/{learner_id}/progress")

    repetitions = scenario["correct_repetitions_for_strong_concept"]
    for _ in range(repetitions):
        strong_quiz_response = await client.post(
            f"/documents/{document_id}/quiz",
            json={"question_count": 10, "learner_id": learner_id},
        )
        if strong_quiz_response.status_code != 200:
            raise RuntimeError("Could not generate simulated strong-concept quiz.")
        strong_quiz = strong_quiz_response.json()
        question = next(
            item
            for item in strong_quiz["questions"]
            if item["concept_label"] == strong_question["concept_label"]
        )
        await client.post(
            weak_submit_path,
            json={
                "quiz_id": strong_quiz["quiz_id"],
                "learner_id": learner_id,
                "answers": [{
                    "question_id": question["question_id"],
                    "selected_option_id": question["answer_key"],
                }],
            },
        )

    progress_response = await client.get(f"/learners/{learner_id}/progress")
    progress = progress_response.json()
    concept_progress = {
        concept["concept_label"]: concept
        for concept in progress["concepts_attempted"]
        if concept["document_id"] == document_id
    }
    weak_progress = concept_progress[weak_question["concept_label"]]
    strong_progress = concept_progress[strong_question["concept_label"]]

    adaptive_response = await client.post(
        f"/learners/{learner_id}/documents/{document_id}/quiz",
        json={"question_count": 10},
    )
    adaptive_questions = adaptive_response.json().get("questions", [])
    adaptive_labels = [question["concept_label"] for question in adaptive_questions]
    adaptive_ids = {question["question_id"] for question in adaptive_questions}
    attempted_ids = {
        question["question_id"]
        for quiz_question in (weak_question, strong_question)
        for question in [quiz_question]
    }
    repeated_adaptive_questions = len(adaptive_ids.intersection(attempted_ids))

    unseen_selected = unseen_label in adaptive_labels
    strong_position = next(
        (index for index, label in enumerate(adaptive_labels) if label == strong_question["concept_label"]),
        None,
    )
    unseen_position = next(
        (index for index, label in enumerate(adaptive_labels) if label == unseen_label),
        None,
    )

    isolation_document_id = document_ids[scenario["isolation_document_id"]]
    isolation_adaptive = await client.post(
        f"/learners/{learner_id}/documents/{isolation_document_id}/quiz",
        json={"question_count": 5},
    )
    isolation_questions = isolation_adaptive.json().get("questions", [])
    isolation_record_ids = set(records_by_document[scenario["isolation_document_id"]])

    other_learner_response = await client.post("/learners")
    other_learner_id = other_learner_response.json()["learner_id"]
    other_progress = await client.get(f"/learners/{other_learner_id}/progress")

    weak_first = bool(adaptive_labels) and adaptive_labels[0] == weak_question["concept_label"]
    strong_not_overprioritized = (
        strong_position is None
        or unseen_position is None
        or strong_position > unseen_position
    )
    return {
        "metrics": {
            "weak_concept_mastery": rate(int(weak_progress["mastery_estimate"] < 0.5), 1),
            "strong_concept_mastery": rate(int(strong_progress["mastery_estimate"] >= 0.7), 1),
            "weak_concept_prioritized": rate(int(weak_first), 1),
            "unseen_concept_selected": rate(int(unseen_selected), 1),
            "strong_concept_not_overprioritized": rate(int(strong_not_overprioritized), 1),
            "retry_not_double_counted": rate(
                int(
                    progress_after_weak.json()["question_attempt_count"]
                    == progress_after_retry.json()["question_attempt_count"]
                    and progress_after_weak.json()["quiz_attempt_count"]
                    == progress_after_retry.json()["quiz_attempt_count"]
                ),
                1,
            ),
            "adaptive_exact_repeat_avoidance": rate(
                len(adaptive_questions) - repeated_adaptive_questions,
                len(adaptive_questions),
            ),
            "learner_isolation": rate(
                int(
                    other_progress.status_code == 200
                    and other_progress.json()["question_attempt_count"] == 0
                    and other_progress.json()["concepts_attempted"] == []
                ),
                1,
            ),
            "document_isolation": rate(
                sum(question.get("source_record_id") in isolation_record_ids for question in isolation_questions),
                len(isolation_questions),
            ),
        },
        "policy_checks": {
            "weak_concept": weak_question["concept_label"],
            "strong_concept": strong_question["concept_label"],
            "unseen_concept": unseen_label,
            "weak_mastery_estimate": weak_progress["mastery_estimate"],
            "strong_mastery_estimate": strong_progress["mastery_estimate"],
            "weak_question_alternates_available": bool(adaptive_labels),
            "repeated_submission_http_status": retry_weak_submit.status_code,
            "adaptive_question_count": len(adaptive_questions),
        },
    }


def _check(name: str, metric: dict[str, Any], condition: bool, expectation: str) -> dict[str, Any]:
    return {
        "name": name,
        "passed": bool(condition),
        "observed": metric["value"],
        "expected": expectation,
    }


async def run_evaluation(dataset: dict[str, Any]) -> dict[str, Any]:
    old_database_path = os.environ.get("NEXUS_DATABASE_PATH")
    with tempfile.TemporaryDirectory(prefix="nexus-task7-eval-") as temporary_directory:
        database_path = Path(temporary_directory) / "evaluation.sqlite3"
        os.environ["NEXUS_DATABASE_PATH"] = str(database_path)
        try:
            if get_database_path() != database_path:
                raise RuntimeError("Evaluation did not use its temporary database.")
            async with httpx.AsyncClient(
                transport=httpx.ASGITransport(app=app), base_url="http://evaluation.local"
            ) as client:
                document_ids, records_by_document = await upload_fixture_documents(client, dataset)
                qa_metrics, _ = await evaluate_qa(
                    client, dataset, document_ids, records_by_document
                )
                quiz_results = await evaluate_quizzes(
                    client, dataset, document_ids, records_by_document
                )
                personalization = await evaluate_personalization(
                    client, dataset, document_ids, records_by_document
                )
        finally:
            if old_database_path is None:
                os.environ.pop("NEXUS_DATABASE_PATH", None)
            else:
                os.environ["NEXUS_DATABASE_PATH"] = old_database_path

    quiz_metrics = quiz_results["metrics"]
    learner_metrics = personalization["metrics"]
    checks = [
        _check("answerable_question_coverage", qa_metrics["answerable_coverage"], qa_metrics["answerable_coverage"]["value"] == 1.0, "1.0"),
        _check("citation_record_accuracy", qa_metrics["citation_source_record_accuracy"], qa_metrics["citation_source_record_accuracy"]["value"] == 1.0, "1.0"),
        _check("citation_location_accuracy", qa_metrics["citation_location_accuracy"], qa_metrics["citation_location_accuracy"]["value"] == 1.0, "1.0"),
        _check("fixture_evidence_support", qa_metrics["evidence_support_rate"], qa_metrics["evidence_support_rate"]["value"] == 1.0, "1.0"),
        _check("unanswerable_abstention", qa_metrics["appropriate_abstention_rate"], qa_metrics["appropriate_abstention_rate"]["value"] == 1.0, "1.0"),
        _check("quiz_source_references", quiz_metrics["valid_source_reference_rate"], quiz_metrics["valid_source_reference_rate"]["value"] == 1.0, "1.0"),
        _check("quiz_answer_keys", quiz_metrics["answer_key_consistency_rate"], quiz_metrics["answer_key_consistency_rate"]["value"] == 1.0, "1.0"),
        _check("quiz_duplicates", quiz_metrics["duplicate_question_rate"], quiz_metrics["duplicate_question_rate"]["value"] == 0.0, "0.0"),
        _check("quiz_url_heading_options", quiz_metrics["url_or_heading_option_rate"], quiz_metrics["url_or_heading_option_rate"]["value"] == 0.0, "0.0"),
        _check("quiz_quality_rules", quiz_metrics["quality_rule_satisfaction_rate"], quiz_metrics["quality_rule_satisfaction_rate"]["value"] == 1.0, "1.0"),
        _check("insufficient_source_reported", quiz_metrics["insufficient_content_report_rate"], quiz_metrics["insufficient_content_report_rate"]["value"] == 1.0, "1.0"),
    ]
    checks.extend(
        _check(name, metric, metric["value"] == 1.0, "1.0")
        for name, metric in learner_metrics.items()
    )
    unsupported_metric = qa_metrics["unsupported_answer_rate"]
    checks.append(
        _check(
            "no_unsupported_answers",
            unsupported_metric,
            unsupported_metric["value"] == 0.0,
            "0.0",
        )
    )

    return {
        "report_version": "1.0.0",
        "dataset_version": dataset["dataset_version"],
        "evaluation_mode": "offline_deterministic_synthetic_fixtures",
        "sample_counts": {
            "documents": len(dataset["documents"]),
            "source_records": sum(len(document["records"]) for document in dataset["documents"]),
            "qa_cases": len(dataset["qa_cases"]),
            "answerable_qa_cases": sum(case["answerable"] for case in dataset["qa_cases"]),
            "unanswerable_qa_cases": sum(not case["answerable"] for case in dataset["qa_cases"]),
            "quiz_checks": len(dataset["quiz_checks"]),
            "simulated_learners": 2,
        },
        "metrics": {
            "grounded_qa": qa_metrics,
            "quiz": quiz_metrics,
            "personalization": learner_metrics,
        },
        "checks": checks,
        "summary": {
            "passed": sum(check["passed"] for check in checks),
            "failed": sum(not check["passed"] for check in checks),
            "total": len(checks),
        },
        "quiz_check_details": quiz_results["quiz_checks"],
        "learner_policy_details": personalization["policy_checks"],
        "limitations": [
            "Evidence support is checked by manually specified fixture terms, not semantic entailment.",
            "Citation correctness and answer support are measured separately.",
            "Keyword overlap can produce unsupported answers for misleading queries; this is measured as a failure, not hidden.",
            "Quiz checks validate deterministic template rules, not pedagogical quality or distractor usefulness.",
            "Learner results are deterministic policy checks over synthetic histories, not evidence of real-world educational effectiveness.",
            "No RAGAS, DeepEval, TruLens, LLM, embedding model, or external service is used.",
        ],
        "reproduction": [
            "uv run --python 3.13 pytest",
            "uv run --python 3.13 python -m evaluation.runner",
        ],
    }


def render_markdown(report: dict[str, Any]) -> str:
    lines = [
        "# NEXUS ACADEMY Offline Evaluation",
        "",
        f"- Dataset version: `{report['dataset_version']}`",
        f"- Mode: `{report['evaluation_mode']}`",
        f"- Checks passed: {report['summary']['passed']}/{report['summary']['total']}",
        "",
        "## Sample Counts",
        "",
        "| Sample | Count |",
        "|---|---:|",
    ]
    lines.extend(
        f"| {sample_name.replace('_', ' ')} | {count} |"
        for sample_name, count in report["sample_counts"].items()
    )
    lines.extend([
        "",
        "All ratios include explicit numerators and denominators. A failed check records current behavior; it is not concealed by the runner.",
        "",
    ])
    for group_name, metrics in report["metrics"].items():
        lines.extend([
            f"## {group_name.replace('_', ' ').title()}",
            "",
            "| Metric | Value | Numerator / denominator |",
            "|---|---:|---:|",
        ])
        for metric_name, metric in metrics.items():
            value = "n/a" if metric["value"] is None else f"{metric['value']:.4f}"
            lines.append(
                f"| {metric_name.replace('_', ' ')} | {value} | "
                f"{metric['numerator']} / {metric['denominator']} |"
            )
        lines.append("")
    lines.extend(["## Checks", "", "| Check | Result | Observed | Expected |", "|---|---|---:|---|"])
    for check in report["checks"]:
        status = "PASS" if check["passed"] else "FAIL"
        observed = "n/a" if check["observed"] is None else str(check["observed"])
        lines.append(f"| {check['name']} | {status} | {observed} | {check['expected']} |")
    lines.extend(["", "## Limitations", ""])
    lines.extend(f"- {limitation}" for limitation in report["limitations"])
    lines.extend(["", "## Reproduction", "", "```powershell"])
    lines.extend(report["reproduction"])
    lines.extend(["```", ""])
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description="Run the NEXUS ACADEMY offline evaluation.")
    parser.add_argument("--dataset", type=Path, default=DEFAULT_DATASET)
    parser.add_argument("--report-directory", type=Path, default=DEFAULT_REPORT_DIRECTORY)
    arguments = parser.parse_args()

    dataset = load_dataset(arguments.dataset)
    report = asyncio.run(run_evaluation(dataset))
    arguments.report_directory.mkdir(parents=True, exist_ok=True)
    json_path = arguments.report_directory / "task7_v1.json"
    markdown_path = arguments.report_directory / "task7_v1.md"
    json_path.write_text(json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    markdown_path.write_text(render_markdown(report), encoding="utf-8")
    print(json.dumps(report, indent=2, sort_keys=True))
    print(f"\nMarkdown report: {markdown_path.relative_to(ROOT)}")
    print(f"JSON report: {json_path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()